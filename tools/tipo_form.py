"""Fill TIPO's official trademark application form (商標註冊申請書 T0101, .docx converted from the official .doc) and export PDF via Word.

The form's check boxes are two-space runs with a character border (w:bdr); "□內以英文字母「v」選填" -> the box run becomes "V ".
Fields marked ※ (組群代碼, 註冊號) are left empty as the form requires. Signatures / seals are left for the applicant.

python tools/tipo_form.py <cases.json>
cases.json: {"template": ".../商標註冊申請書_official.docx", "out_dir": "...", "cases": [ {...}, ... ]}
"""
import copy
import json
import subprocess
import sys
from pathlib import Path

import docx


def is_box(run):
    rpr = run._r.rPr
    return rpr is not None and rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bdr") is not None and run.text.strip() == ""


def tick(par, label):
    """Put V in the bordered box run that precedes the run whose text starts with label."""
    runs = par.runs
    for i, r in enumerate(runs):
        if r.text.strip().startswith(label):
            for j in range(i - 1, -1, -1):
                if is_box(runs[j]):
                    runs[j].text = "V "
                    return True
    return False


def append(par, text):
    r = par.add_run(text)
    if par.runs and len(par.runs) > 1:                      # inherit the label's font (標楷體)
        src = par.runs[-2]._r.rPr
        if src is not None:
            r._r.insert(0, copy.deepcopy(src))
            b = r._r.rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bdr")
            if b is not None:
                r._r.rPr.remove(b)
    return r


def cells_unique(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            out.append(c)
    return out


def replace_in(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    full = "".join(r.text for r in par.runs)
    if old in full:                                          # spread across runs: rewrite the first run, clear the others
        par.runs[0].text = full.replace(old, new, 1)
        for r in par.runs[1:]:
            r.text = ""
        return True
    return False


def fill(template, case, out_docx):
    d = docx.Document(template)
    P = d.paragraphs
    find = lambda s: next(p for p in P if p.text.strip().startswith(s))      # noqa: E731
    replace_in(find("◎繳納商標規費金額"), "          元", f"  {case['fee']}  元")
    append(find("一、商標名稱："), case["name"])
    tick(find("二、商標圖樣顏色："), case.get("color", "墨色"))
    if case.get("disclaim"):
        append(find("三、聲明不專用："), case["disclaim"])
    for k in ("中文", "外文", "語文別", "中文字義", "圖形", "記號"):
        v = case["analysis"].get(k, "")
        if v:
            append(find(k + "："), v)
    replace_in(find("肆、申請人"), "共   人", "共 1 人")
    A = case["applicant"]
    t = d.tables[0]
    rows = t.rows
    nat = cells_unique(rows[1])[-1]
    tick(nat.paragraphs[0], "中華民國")
    kind = cells_unique(rows[2])[-1]
    for p in kind.paragraphs:
        tick(p, A["kind"])
    def cell_with(row, marker):                              # the cell that holds the label (merged cells repeat; take the first match)
        for c in cells_unique(row):
            if marker in c.text:
                return c
        return cells_unique(row)[-1]
    for ri, marker, key in ((3, "ID", "id"), (4, "(中文)", "name_zh"), (5, "(英文)", "name_en"), (6, "(中文)", "rep_zh"), (7, "(英文)", "rep_en"),
                            (8, "(中文)", "addr_zh"), (9, "(英文)", "addr_en"), (11, "聯絡電話", "tel"), (12, "真", "fax"), (13, "E-MAIL", "email")):
        if A.get(key):
            append(cell_with(rows[ri], marker).paragraphs[0], " " + A[key])
    # classes: three 類別 / 商品／服務名稱 blocks in the form
    classes = case["classes"]
    replace_in(find("指定申請類別：第"), "類別：第", "類別：第 " + "、".join(c for c, _ in classes) + " ")
    blocks = [i for i, p in enumerate(P) if p.text.strip() == "類別："]
    for (cls, names), bi in zip(classes, blocks):
        append(P[bi], cls)
        append(P[bi + 1], "")
        P[bi + 2].add_run("；".join(names) + "。")
    extra = classes[len(blocks):]
    if extra:                                                 # more classes than blocks: continue after the last block (form allows 附表)
        note = P[blocks[-1] + 2]
        for cls, names in extra:
            note.add_run(f"\n類別：{cls}\n商品／服務名稱：" + "；".join(names) + "。")
    # 具結: applicant without agent -> first statement
    tick(next(p for p in P if "本申請書所填寫之資料係為真實" in p.text), "本申請書所填寫之資料係為真實")
    tick(next(p for p in P if "商標圖樣浮貼一式" in p.text), "商標圖樣浮貼一式")
    d.save(out_docx)


def to_pdf(docx_path, pdf_path):
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; $d = $w.Documents.Open('{docx_path}'); "
          f"$d.ExportAsFixedFormat('{pdf_path}', 17); $d.Close(0); $w.Quit(); 'ok'")
    return subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True).stdout.strip()


if __name__ == "__main__":
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out = Path(spec["out_dir"])
    out.mkdir(parents=True, exist_ok=True)
    for c in spec["cases"]:
        dx = out / f"{c['file']}.docx"
        fill(spec["template"], c, str(dx))
        print(c["file"], to_pdf(str(dx), str(out / f"{c['file']}.pdf")))
